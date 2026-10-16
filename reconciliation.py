import math
import os
import re
import tempfile
from pathlib import Path

import openpyxl
import pandas as pd
from openpyxl.utils.dataframe import dataframe_to_rows


SOA_REQUIRED_COLUMNS = {"Invoice Number", "Amount( Original )"}
JDE_REQUIRED_COLUMNS = {"Gross Amount", "Reg Code", "Pay Stat"}
CORA_REQUIRED_COLUMNS = {"Invoice No", "Current State", "Total Amount"}


class ReconciliationValidationError(ValueError):
    """Raised when a reconciliation input is unsafe or has an invalid schema."""


def normalize_key(val):
    if pd.isna(val) or val is None:
        return ""
    s = str(val).strip()
    if s.endswith(".0"):
        s = s[:-2]
    return s


def clean_alphanumeric(val):
    s = normalize_key(val)
    for ch in ["-", "/", " ", "_", ".", "#", ":"]:
        s = s.replace(ch, "")
    return s.lower()


def format_date_ddmmyyyy(val):
    if pd.isna(val) or val is None or str(val).strip() in ["", "NaT", "nan", "None"]:
        return ""
    try:
        if hasattr(val, "strftime"):
            return val.strftime("%d/%m/%Y")
        dt = pd.to_datetime(val, dayfirst=True)
        return dt.strftime("%d/%m/%Y")
    except Exception:
        return str(val).strip()


def _to_datetime(val):
    if pd.isna(val) or val is None or str(val).strip() in ["", "NaT", "nan", "None"]:
        return None
    try:
        return val if hasattr(val, "strftime") else pd.to_datetime(val, dayfirst=True)
    except Exception:
        return None


def payment_cycle_status(due_date, as_of):
    """Name the payment cycle; an overdue held invoice goes in the current cycle."""
    dt = _to_datetime(due_date)
    if dt is None:
        return "Ready for Payment and scheduled for Payment Cycle"
    cycle = (dt.year, dt.month)
    if cycle < (as_of.year, as_of.month):
        cycle = (as_of.year, as_of.month)
    month_name = pd.Timestamp(year=cycle[0], month=cycle[1], day=1).strftime("%B")
    if cycle[0] != as_of.year:
        month_name += f"'{cycle[0] % 100:02d}"
    return f"Ready for Payment and scheduled for {month_name} Payment Cycle"


def parse_amount(value, *, source, row_number, required=False):
    """Parse a UK-style amount and reject malformed non-empty values."""
    if value is None or pd.isna(value) or str(value).strip() == "":
        if required:
            raise ReconciliationValidationError(
                f"{source} row {row_number}: amount is required."
            )
        return None

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return round(float(value), 2)

    raw = str(value).strip()
    negative = raw.startswith("(") and raw.endswith(")")
    if negative:
        raw = raw[1:-1].strip()

    cleaned = re.sub(r"(?i)\b(?:GBP|USD|EUR)\b", "", raw)
    for symbol in ("£", "$", "€", ",", " ", "\u00a0"):
        cleaned = cleaned.replace(symbol, "")

    if not re.fullmatch(r"[+-]?(?:\d+(?:\.\d+)?|\.\d+)", cleaned):
        raise ReconciliationValidationError(
            f"{source} row {row_number}: invalid amount {value!r}."
        )

    parsed = float(cleaned)
    return round(-parsed if negative else parsed, 2)


def _require_columns(df, required, source):
    missing = sorted(required.difference(df.columns))
    if missing:
        raise ReconciliationValidationError(
            f"{source} is missing required column(s): {', '.join(missing)}."
        )
    if df.empty:
        raise ReconciliationValidationError(f"{source} contains no data rows.")


def validate_soa_dataframe(df_soa):
    if not isinstance(df_soa, pd.DataFrame):
        raise ReconciliationValidationError("SOA input must be an Excel path or DataFrame.")
    _require_columns(df_soa, SOA_REQUIRED_COLUMNS, "SOA data")

    for offset, (_, row) in enumerate(df_soa.iterrows(), start=2):
        invoice_number = normalize_key(row.get("Invoice Number"))
        if not invoice_number:
            raise ReconciliationValidationError(
                f"SOA data row {offset}: Invoice Number is required."
            )
        parse_amount(
            row.get("Amount( Original )"),
            source="SOA data",
            row_number=offset,
        )
        parse_amount(
            row.get("Open amount"),
            source="SOA data Open amount",
            row_number=offset,
        )


def validate_jde_dataframe(df_jde):
    _require_columns(df_jde, JDE_REQUIRED_COLUMNS, "JDE report")
    if "Invoice Number" not in df_jde.columns and "Document Number" not in df_jde.columns:
        raise ReconciliationValidationError(
            "JDE report requires Invoice Number or Document Number."
        )

    for offset, (_, row) in enumerate(df_jde.iterrows(), start=2):
        if not normalize_key(row.get("Invoice Number")) and not normalize_key(row.get("Document Number")):
            continue
        parse_amount(
            row.get("Gross Amount"),
            source="JDE report Gross Amount",
            row_number=offset,
        )


def validate_cora_dataframe(df_cora):
    _require_columns(df_cora, CORA_REQUIRED_COLUMNS, "Cora report")
    for offset, (_, row) in enumerate(df_cora.iterrows(), start=2):
        if not normalize_key(row.get("Invoice No")):
            continue
        parse_amount(
            row.get("Total Amount"),
            source="Cora report Total Amount",
            row_number=offset,
        )


def _read_excel_sheet(path, preferred_sheet, source):
    if not path or not os.path.isfile(path):
        raise ReconciliationValidationError(f"{source} file does not exist: {path}")
    try:
        with pd.ExcelFile(path) as workbook:
            sheet_name = preferred_sheet if preferred_sheet in workbook.sheet_names else workbook.sheet_names[0]
            return pd.read_excel(workbook, sheet_name=sheet_name)
    except ReconciliationValidationError:
        raise
    except Exception as exc:
        raise ReconciliationValidationError(f"Could not read {source}: {exc}") from exc


def validate_reference_files(cora_file, jde_file):
    """Validate both source workbooks before any paid AI extraction begins."""
    df_cora = _read_excel_sheet(cora_file, "Cora Report", "Cora report")
    df_jde = _read_excel_sheet(jde_file, "JDE Report", "JDE report")
    validate_cora_dataframe(df_cora)
    validate_jde_dataframe(df_jde)
    return df_cora, df_jde


def _safe_sheet_name(name, default_name="Sheet"):
    safe_name = name or default_name
    for ch in ["*", ":", "?", "[", "]", "/", "\\"]:
        safe_name = safe_name.replace(ch, "")
    return safe_name[:31] if safe_name else default_name


def _amount_matches(left, right):
    return left is not None and right is not None and round(left, 2) == round(right, 2)


def _format_amount(value):
    return "unavailable" if value is None else f"{value:,.2f}"


UNDER_PROCESSING = (
    "Under Processing and will be reviewed by the team, should more information be required we will be in touch"
)
MISSING_IN_LEDGER = "Missing in ledger please submit your invoice to APukinvoices@otis.com"

# AP Advance priority from the VSR "AP Advance & Logic" sheet. When one invoice
# has several Cora rows, the state listed first decides the status. The states
# already reflected in JDE (Paid ... Ready For QC) come before priority 1.
CORA_STATE_PRIORITY = (
    "paid",
    "partially paid",
    "awaiting payment feed",
    "ready for qc",
    "awaiting posting acknowledgement",   # 1
    "awaiting gr",                        # 2
    "awaiting approval",                  # 3
    "ready for gr posting",               # 4
    "ready for gr creation",              # 5
    "awaiting variance or fmi approval",  # 6
    "ready for business help",            # 7
    "awaiting mdm team response",         # 8
    "invoice posting error",              # 9
    "ready for coding by ap",             # 10
    "ready for manual processing",        # 11
    "pending for cora validation",        # 12
    "awaiting ocr response",              # 13
    "returned",                           # 14
    "cancelled",                          # 15
    "ready for archival",                 # 16
    "archived",                           # 17
    "reversed",                           # 18
)
_CORA_RANK = {state: rank for rank, state in enumerate(CORA_STATE_PRIORITY)}

READY_FOR_PAYMENT_STATES = {"awaiting payment feed", "ready for qc"}
PENDING_APPROVAL_REASONS = {
    "awaiting approval": "Approval Required",
    "ready for gr creation": "GR creation",
    "awaiting variance or fmi approval": "Variance or FMI Approval",
    "ready for coding by ap": "Coding by AP",
}
UNDER_PROCESSING_STATES = {
    "awaiting posting acknowledgement",
    "awaiting gr",
    "ready for gr posting",
    "ready for business help",
    "awaiting mdm team response",
    "invoice posting error",
    "ready for manual processing",
    "pending for cora validation",
    "awaiting ocr response",
    "reversed",
    # Not in the VSR sheet, but internal processing steps.
    "under processing",
    "ready for posting",
    "qc corrections",
}


def _cora_sort_key(candidate, amount_soa):
    state = candidate["Current State"].lower()
    return (
        _CORA_RANK.get(state, len(CORA_STATE_PRIORITY)),
        not _amount_matches(candidate["Total Amount"], amount_soa),
    )


def _reason(*values):
    """First meaningful reason; the generic 'Others' category is skipped."""
    for value in values:
        text = (value or "").strip().rstrip(".").strip()
        if text and text.lower() != "others":
            return text
    return ""


def _rejected_status(matched_cora):
    reason = _reason(matched_cora["Return Comment"], matched_cora["Return Reason"])
    status = f"Rejected due to {reason or 'Reason Mentioned in AP Advance'}"
    owner = matched_cora["Cora Current Owner"]
    return f"{status} by {owner}" if owner and owner.lower() != "cora" else status


def _cancelled_status(matched_cora):
    reason = _reason(matched_cora["Cancel Reason"], matched_cora["Return Comment"])
    return f"Cancelled due to {reason or 'Reason Mentioned in AP Advance'}"


def _jde_status(matched_jde, cora_payment_date, as_of):
    """JDE decides first: invoices present in JDE are paid or awaiting payment."""
    if not matched_jde:
        return "", ""

    lines = [(reg, stat) for reg, stat in matched_jde["Lines"] if stat]
    if not lines:
        return "", ""
    paid = [stat == "P" for _, stat in lines]
    held = any(stat == "H" and reg in ["L", ""] for reg, stat in lines)

    if all(paid):
        payment_date = format_date_ddmmyyyy(cora_payment_date)
        return (f"Paid in full on {payment_date}" if payment_date else "Paid in full"), ""
    if held and any(paid):
        return "Partially Paid", ""
    if held:
        return payment_cycle_status(matched_jde["Due Date"], as_of), ""
    return "", ""


def _cora_status(matched_cora, as_of):
    if not matched_cora or not matched_cora["Current State"]:
        return "", ""

    cora_state_lower = matched_cora["Current State"].lower()
    cora_owner = matched_cora["Cora Current Owner"]

    if cora_state_lower == "paid":
        payment_date = format_date_ddmmyyyy(matched_cora.get("Payment Date"))
        return (f"Paid in full on {payment_date}" if payment_date else "Paid in full"), ""

    if cora_state_lower == "partially paid":
        payment_date = format_date_ddmmyyyy(matched_cora.get("Payment Date"))
        return (
            f"Partially Paid on {payment_date}" if payment_date else "Partially Paid",
            cora_owner,
        )

    if cora_state_lower in READY_FOR_PAYMENT_STATES:
        # Blank Excel cells arrive as NaN, which is truthy, so test real dates.
        due_date = next(
            (
                matched_cora.get(column) for column in ("Pay By Date", "Payment Date", "Invoice Date")
                if _to_datetime(matched_cora.get(column)) is not None
            ),
            None,
        )
        return payment_cycle_status(due_date, as_of), ""

    if cora_state_lower in PENDING_APPROVAL_REASONS:
        buyer = cora_owner if cora_owner and cora_owner.lower() != "cora" else "Approver"
        reason = PENDING_APPROVAL_REASONS[cora_state_lower]
        return f"Pending Approval from {buyer} due to {reason}", buyer

    if cora_state_lower in UNDER_PROCESSING_STATES:
        pending = cora_owner if cora_owner and cora_owner.lower() != "cora" else ""
        return UNDER_PROCESSING, pending

    if cora_state_lower == "returned":
        return _rejected_status(matched_cora), cora_owner

    if cora_state_lower == "cancelled":
        return _cancelled_status(matched_cora), cora_owner

    if cora_state_lower in ["ready for archival", "archived"]:
        if _reason(matched_cora["Return Comment"], matched_cora["Return Reason"]) or not matched_cora["Cancel Reason"]:
            return _rejected_status(matched_cora), cora_owner
        return _cancelled_status(matched_cora), cora_owner

    if cora_state_lower == "awaiting credit memo":
        return (
            "Need Manual Intervention - Please share updated credit memo with invoice number",
            cora_owner,
        )

    return "", ""


def _amount_check(amount_soa, jde_amount, cora_amount):
    """Amount differences are reported beside the status rather than replacing it."""
    if amount_soa is None:
        return "SOA amount missing"
    differences = [
        f"{label} {_format_amount(value)}"
        for label, value in (("JDE", jde_amount), ("Cora", cora_amount))
        if value is not None and not _amount_matches(value, amount_soa)
    ]
    if not differences:
        return ""
    return f"SOA {_format_amount(amount_soa)} differs from " + ", ".join(differences)


def _validate_output_path(output_file, source_paths):
    if not output_file:
        raise ReconciliationValidationError("An output workbook path is required.")
    output_path = Path(output_file).expanduser().resolve()
    if output_path.suffix.lower() != ".xlsx":
        raise ReconciliationValidationError("The output workbook must use the .xlsx extension.")
    if not output_path.parent.is_dir():
        raise ReconciliationValidationError(
            f"Output directory does not exist: {output_path.parent}"
        )
    for source_path in source_paths:
        if source_path and Path(source_path).expanduser().resolve() == output_path:
            raise ReconciliationValidationError(
                "The output workbook cannot overwrite an input workbook."
            )
    return output_path


def _save_workbook_atomic(workbook, output_path):
    """Write beside the destination, then replace it only after a valid save."""
    temp_path = None
    try:
        with tempfile.NamedTemporaryFile(
            prefix=".reconciliation-", suffix=".xlsx", dir=output_path.parent, delete=False
        ) as temp_file:
            temp_path = temp_file.name
        workbook.save(temp_path)
        os.replace(temp_path, output_path)
        temp_path = None
    finally:
        if temp_path and os.path.exists(temp_path):
            os.remove(temp_path)


def create_reconciliation_file(df_soa, cora_file, jde_file, logic_file, output_file, logger=print, as_of=None):
    """Create the final workbook while preserving the established source matching order."""
    logger("Creating new reconciliation file...")
    as_of = pd.Timestamp(as_of or pd.Timestamp.today()).normalize()

    soa_source_path = None
    if isinstance(df_soa, (str, os.PathLike)):
        soa_source_path = os.fspath(df_soa)
        df_soa = _read_excel_sheet(soa_source_path, "SOA_Status", "SOA input")
        logger(
            f"Loaded {len(df_soa)} SOA records from Excel: {os.path.basename(soa_source_path)}"
        )

    validate_soa_dataframe(df_soa)
    output_path = _validate_output_path(
        output_file, [soa_source_path, cora_file, jde_file, logic_file]
    )

    # JDE may contain several accounting lines for one invoice. Build one
    # candidate per invoice using the sum of all its Gross Amount lines, then
    # retain the established exact/zero-stripped/alphanumeric lookup order.
    logger(f"Loading JDE Report data from {jde_file}...")
    df_jde = _read_excel_sheet(jde_file, "JDE Report", "JDE report")
    validate_jde_dataframe(df_jde)
    jde_exact_map = {}
    jde_alpha_map = {}
    jde_groups = {}

    for row_number, (_, row) in enumerate(df_jde.iterrows(), start=2):
        inv_raw = row.get("Invoice Number")
        doc_raw = row.get("Document Number")
        gross_amt = parse_amount(
            row.get("Gross Amount"), source="JDE report Gross Amount", row_number=row_number
        )
        reg_code = str(row.get("Reg Code", "")).strip().upper()
        if reg_code in ["NAN", "NONE"]:
            reg_code = ""
        pay_stat = str(row.get("Pay Stat", "")).strip().upper()
        if pay_stat in ["NAN", "NONE"]:
            pay_stat = ""
        due_date = row.get("Discount Due Date")
        if pd.isna(due_date):
            due_date = row.get("Net Due Date")

        rec_data = {
            "Invoice Number": str(inv_raw) if pd.notna(inv_raw) else "",
            "Document Number": str(doc_raw) if pd.notna(doc_raw) else "",
            "Reg Code": reg_code,
            "Pay Stat": pay_stat,
            "Due Date": due_date,
            "Gross Amount": gross_amt,
            "Payment Date": row.get("Payment Date"),
            "GL Date": row.get("G/L Date"),
            "Batch Date": row.get("Batch Date"),
            "Doc Type": row.get("Doc Type"),
        }

        identifiers = [
            raw for raw in (inv_raw, doc_raw)
            if pd.notna(raw) and normalize_key(raw)
        ]
        if not identifiers:
            continue

        primary_key = normalize_key(identifiers[0]).lower()
        group_key = primary_key.lstrip("0") or primary_key
        group = jde_groups.get(group_key)
        if group is None:
            group = {
                "record": rec_data,
                "amounts": [],
                "exact_keys": set(),
                "alpha_keys": set(),
                "line_count": 0,
            }
            rec_data["Lines"] = []
            jde_groups[group_key] = group

        group["line_count"] += 1
        group["record"]["Lines"].append((reg_code, pay_stat))
        if gross_amt is not None:
            group["amounts"].append(gross_amt)

        for identifier in identifiers:
            normalized = normalize_key(identifier).lower()
            stripped = normalized.lstrip("0")
            alphanumeric = clean_alphanumeric(identifier)
            group["exact_keys"].update(key for key in (normalized, stripped) if key)
            if alphanumeric:
                group["alpha_keys"].add(alphanumeric)

    for group in jde_groups.values():
        rec_data = group["record"]
        rec_data["Gross Amount"] = (
            round(math.fsum(group["amounts"]), 2) if group["amounts"] else None
        )
        rec_data["JDE Line Count"] = group["line_count"]
        for key in group["exact_keys"]:
            jde_exact_map.setdefault(key, []).append(rec_data)
        for key in group["alpha_keys"]:
            jde_alpha_map.setdefault(key, []).append(rec_data)

    logger(
        f" -> Aggregated {len(df_jde)} JDE rows into "
        f"{len(jde_groups)} invoice total(s)."
    )

    logger(f"Loading Cora Report data from {cora_file}...")
    df_cora = _read_excel_sheet(cora_file, "Cora Report", "Cora report")
    validate_cora_dataframe(df_cora)
    cora_exact_map = {}
    cora_alpha_map = {}

    for row_number, (_, row) in enumerate(df_cora.iterrows(), start=2):
        inv_raw = row.get("Invoice No")
        if pd.isna(inv_raw) or not str(inv_raw).strip():
            continue
        cora_state = str(row.get("Current State", "")).strip()
        if cora_state.upper() in ["NAN", "NONE", "NULL"]:
            cora_state = ""
        cora_owner = str(row.get("Cora Current Owner", "")).strip()
        if cora_owner.upper() in ["NAN", "NONE", "NULL"]:
            cora_owner = ""

        total_amount = parse_amount(
            row.get("Total Amount"), source="Cora report Total Amount", row_number=row_number
        )
        reasons = {
            column: str(row.get(column)).strip() if pd.notna(row.get(column)) else ""
            for column in ("Return Reason comments", "Return Reason", "Cancel Reason")
        }
        rec_data = {
            "Invoice No": str(inv_raw).strip(),
            "Current State": cora_state,
            "Cora Current Owner": cora_owner,
            "Total Amount": total_amount,
            "Payment Date": row.get("Payment Date"),
            "Pay By Date": row.get("Pay By Date"),
            "Invoice Date": row.get("Invoice Date"),
            "Return Comment": reasons["Return Reason comments"],
            "Return Reason": reasons["Return Reason"],
            "Cancel Reason": reasons["Cancel Reason"],
            "Voucher Date": row.get("Voucher Creation Date"),
        }

        k_norm = normalize_key(inv_raw).lower()
        k_lstrip = k_norm.lstrip("0")
        k_alpha = clean_alphanumeric(inv_raw)
        for key in {k_norm, k_lstrip}:
            if key:
                cora_exact_map.setdefault(key, []).append(rec_data)
        if k_alpha:
            cora_alpha_map.setdefault(k_alpha, []).append(rec_data)

    logger(f" -> Indexed {len(df_cora)} Cora rows successfully.")
    logger("Performing lookups and applying reconciliation comments...")

    reconciled_rows = []
    for row_number, (_, row) in enumerate(df_soa.iterrows(), start=2):
        inv_soa = normalize_key(row.get("Invoice Number"))
        amount_soa = parse_amount(
            row.get("Amount( Original )"), source="SOA data", row_number=row_number
        )
        open_amount = parse_amount(
            row.get("Open amount"), source="SOA data Open amount", row_number=row_number
        )
        date_soa = format_date_ddmmyyyy(row.get("Date"))
        source_file = row.get("Source File", "")

        k_norm = inv_soa.lower()
        k_lstrip = k_norm.lstrip("0")
        k_alpha = clean_alphanumeric(inv_soa)

        jde_candidates = jde_exact_map.get(k_norm) or jde_exact_map.get(k_lstrip) or jde_alpha_map.get(k_alpha) or []
        matched_jde = None
        if jde_candidates:
            for candidate in jde_candidates:
                if _amount_matches(candidate["Gross Amount"], amount_soa):
                    matched_jde = candidate
                    break
            if not matched_jde:
                for candidate in jde_candidates:
                    if candidate["Gross Amount"] is not None and candidate["Reg Code"] in ["L", "P"]:
                        matched_jde = candidate
                        break
            if not matched_jde:
                matched_jde = jde_candidates[0]

        # One invoice can have several Cora rows (resubmissions, duplicates).
        # The AP Advance priority decides which state is reported.
        cora_candidates = cora_exact_map.get(k_norm) or cora_exact_map.get(k_lstrip) or cora_alpha_map.get(k_alpha) or []
        matched_cora = (
            min(cora_candidates, key=lambda candidate: _cora_sort_key(candidate, amount_soa))
            if cora_candidates else None
        )
        cora_payment_date = next(
            (
                candidate["Payment Date"] for candidate in cora_candidates
                if candidate["Current State"].lower() == "paid" and _to_datetime(candidate["Payment Date"]) is not None
            ),
            None,
        )

        reg_code = matched_jde["Reg Code"] if matched_jde else ""
        pay_stat = matched_jde["Pay Stat"] if matched_jde else ""
        due_date = matched_jde["Due Date"] if matched_jde else ""
        jde_amount = matched_jde["Gross Amount"] if matched_jde else None
        cora_state = matched_cora["Current State"] if matched_cora else ""
        cora_owner = matched_cora["Cora Current Owner"] if matched_cora else ""
        cora_amount = matched_cora["Total Amount"] if matched_cora else None

        # JDE (paid / ready for payment) takes precedence over Cora. A small
        # amount difference no longer hides the status; it is shown in
        # "Amount Check" instead.
        jde_status, jde_pending = _jde_status(matched_jde, cora_payment_date, as_of)
        cora_status, cora_pending = _cora_status(matched_cora, as_of)

        if jde_status:
            status = jde_status
            pending_with = jde_pending
        elif cora_status:
            status = cora_status
            pending_with = cora_pending
        elif matched_jde:
            # In JDE but on hold outside the payment run (for example Reg Code P).
            status = UNDER_PROCESSING
            pending_with = ""
        elif matched_cora:
            status = "Matched - Manual review required"
            pending_with = cora_owner
        else:
            status = MISSING_IN_LEDGER
            pending_with = ""

        reconciled_rows.append({
            "Date": date_soa,
            "Invoice Number": inv_soa,
            "Amount( Original )": amount_soa,
            "Open amount": open_amount,
            "Invoice Status": status,
            "Pending with": pending_with,
            "Reg Code": reg_code,
            "Pay Status": pay_stat,
            "Due Date": format_date_ddmmyyyy(due_date),
            "Cora Current State": cora_state,
            "Cora Current Owner": cora_owner,
            "Amount Check": _amount_check(amount_soa, jde_amount, cora_amount),
            "Source File": source_file,
        })

    df_reconciled = pd.DataFrame(reconciled_rows)
    workbook = openpyxl.Workbook()
    workbook.remove(workbook.active)

    soa_sheet = workbook.create_sheet("SOA_Status", 0)
    for row in dataframe_to_rows(df_reconciled, index=False, header=True):
        soa_sheet.append(row)
    logger(f" -> Added 'SOA_Status' with {len(df_reconciled)} reconciled rows.")

    cora_sheet = workbook.create_sheet("Cora")
    for row in dataframe_to_rows(df_cora, index=False, header=True):
        cora_sheet.append(row)

    jde_sheet = workbook.create_sheet("Supplier Ledger")
    for row in dataframe_to_rows(df_jde, index=False, header=True):
        jde_sheet.append(row)

    if logic_file:
        if not os.path.isfile(logic_file):
            raise ReconciliationValidationError(f"Logic workbook does not exist: {logic_file}")
        try:
            with pd.ExcelFile(logic_file) as logic_workbook:
                for sheet_name in logic_workbook.sheet_names:
                    logic_df = pd.read_excel(logic_workbook, sheet_name=sheet_name)
                    output_sheet = workbook.create_sheet(_safe_sheet_name(sheet_name, sheet_name))
                    for row in dataframe_to_rows(logic_df, index=False, header=True):
                        output_sheet.append(row)
        except Exception as exc:
            raise ReconciliationValidationError(f"Could not copy the logic workbook: {exc}") from exc

    _save_workbook_atomic(workbook, output_path)

    logger(f"SUCCESS! Created final reconciliation file: {output_path}")
    logger(f"Total Sheets ({len(workbook.sheetnames)}): {workbook.sheetnames}")
    return df_reconciled


if __name__ == "__main__":
    create_reconciliation_file(
        "Output_SOA_Status.xlsx",
        "Cora-Report.xlsx",
        "JDE Report.xlsx",
        "Core-Logic.xlsx",
        "Final_Reconciliation_Output.xlsx",
    )
