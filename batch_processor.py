import os
import email
import tempfile
from email import policy
import pandas as pd
from docx import Document
from soa_extraction import (
    extract_msg_contents,
    call_llm_for_extraction,
    ExtractedSOA,
    DIRECT_MIME_TYPES
)

def extract_eml_contents(file_path):
    """
    Parses a .eml file and extracts its HTML body and attachments into the unified schema.
    """
    with open(file_path, 'rb') as f:
        msg = email.message_from_binary_file(f, policy=policy.default)
        
    subject = msg.get('subject', 'No Subject')
    sender = msg.get('from', 'Unknown Sender')
    date = msg.get('date', 'Unknown Date')
    
    # Extract body (prefer HTML)
    body_html = ""
    body_part = msg.get_body(preferencelist=('html', 'plain'))
    if body_part:
        try:
            body_html = body_part.get_content()
        except Exception:
            body_html = str(body_part.get_payload(decode=True), 'utf-8', 'ignore')
            
    attachments_extracted = []
    
    # Extract attachments
    for part in msg.iter_attachments():
        att_name = part.get_filename()
        if not att_name:
            continue
            
        content_bytes = part.get_content()
        if isinstance(content_bytes, str):
            content_bytes = content_bytes.encode('utf-8')
            
        _, ext = os.path.splitext(att_name.lower())
        
        attachment_info = {
            "filename": att_name,
            "extension": ext,
            "bytes_length": len(content_bytes)
        }
        
        if ext in ['.xlsx', '.xls', '.xlsm']:
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as temp_file:
                temp_file.write(content_bytes)
                temp_path = temp_file.name
            try:
                excel_file = pd.ExcelFile(temp_path)
                csv_content = ""
                for sheet_name in excel_file.sheet_names:
                    df = pd.read_excel(temp_path, sheet_name=sheet_name)
                    df_cleaned = df.copy()
                    for col in df_cleaned.columns:
                        if pd.api.types.is_datetime64_any_dtype(df_cleaned[col]):
                            df_cleaned[col] = df_cleaned[col].dt.strftime('%Y-%m-%d %H:%M:%S')
                    csv_content += f"--- SHEET: {sheet_name} ---\n"
                    csv_content += df_cleaned.to_csv(index=False) + "\n\n"
                attachment_info["mime_type"] = "text/csv"
                attachment_info["raw_bytes"] = csv_content.encode("utf-8")
                attachment_info["send_via_bytes"] = True
                excel_file.close()
            except Exception as e:
                print(f"   Error parsing Excel file {att_name}: {e}")
            finally:
                if os.path.exists(temp_path):
                    try:
                        os.remove(temp_path)
                    except Exception:
                        pass
                        
        elif ext in DIRECT_MIME_TYPES:
            attachment_info["mime_type"] = DIRECT_MIME_TYPES[ext]
            attachment_info["raw_bytes"] = content_bytes
            attachment_info["send_via_bytes"] = True
            
        elif ext == '.docx':
            with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as temp_file:
                temp_file.write(content_bytes)
                temp_path = temp_file.name
            try:
                doc = Document(temp_path)
                fullText = []
                for para in doc.paragraphs:
                    fullText.append(para.text)
                for table in doc.tables:
                    for row in table.rows:
                        row_text = [cell.text for cell in row.cells]
                        fullText.append(" | ".join(row_text))
                attachment_info["word_text"] = "\n".join(fullText)
            except Exception as e:
                print(f"   Error parsing Word Document {att_name}: {e}")
            finally:
                if os.path.exists(temp_path):
                    try:
                        os.remove(temp_path)
                    except Exception:
                        pass
                        
        attachments_extracted.append(attachment_info)
        
    return {
        "subject": subject,
        "sender": sender,
        "date": date,
        "body_html": body_html,
        "attachments": attachments_extracted
    }

def extract_standalone_file_contents(file_path):
    """
    Wraps a standalone file (.pdf, .xlsx, .xls, .xlsm, .csv, .docx, images, etc.) into the same 
    payload structure that the MSG extractor creates so the LLM can process it.
    """
    filename = os.path.basename(file_path)
    _, ext = os.path.splitext(filename.lower())
    
    with open(file_path, "rb") as f:
        content_bytes = f.read()
        
    attachment_info = {
        "filename": filename,
        "extension": ext,
        "bytes_length": len(content_bytes)
    }
    
    # 1. Excel Handling
    if ext in ['.xlsx', '.xls', '.xlsm']:
        excel_file = pd.ExcelFile(file_path)
        csv_content = ""
        for sheet_name in excel_file.sheet_names:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
            df_cleaned = df.copy()
            for col in df_cleaned.columns:
                if pd.api.types.is_datetime64_any_dtype(df_cleaned[col]):
                    df_cleaned[col] = df_cleaned[col].dt.strftime('%Y-%m-%d %H:%M:%S')
            csv_content += f"--- SHEET: {sheet_name} ---\n"
            csv_content += df_cleaned.to_csv(index=False) + "\n\n"
            
        attachment_info["mime_type"] = "text/csv"
        attachment_info["raw_bytes"] = csv_content.encode('utf-8')
        attachment_info["send_via_bytes"] = True
        excel_file.close() # Keep windows file handles happy
        
    # 2. Native Types (PDF, Images, Text, etc.)
    elif ext in DIRECT_MIME_TYPES:
        attachment_info["mime_type"] = DIRECT_MIME_TYPES[ext]
        attachment_info["raw_bytes"] = content_bytes
        attachment_info["send_via_bytes"] = True
        
    # 3. Word Documents
    elif ext == '.docx':
        doc = Document(file_path)
        fullText = []
        for para in doc.paragraphs:
            fullText.append(para.text)
        for table in doc.tables:
            for row in table.rows:
                row_text = [cell.text for cell in row.cells]
                fullText.append(" | ".join(row_text))
        attachment_info["word_text"] = "\n".join(fullText)
        
    return {
        "subject": f"Standalone File Processing: {filename}",
        "sender": "Direct Upload",
        "date": "N/A",
        "body_html": f"<p>Please extract Statement of Account (SOA) line items from the attached document: <b>{filename}</b></p>",
        "attachments": [attachment_info]
    }

def batch_process_folder(folder_path, output_excel_path):
    all_records = []
    
    # Iterate through every file in the folder
    for filename in os.listdir(folder_path):
        file_path = os.path.join(folder_path, filename)
        if not os.path.isfile(file_path):
            continue
            
        print(f"\n==================================================")
        print(f"Processing File: {filename}")
        print(f"==================================================")
        
        _, ext = os.path.splitext(filename.lower())
        
        try:
            # Route to appropriate parser for emails or standalone files
            if ext == '.msg':
                parsed_data = extract_msg_contents(file_path)
            elif ext == '.eml':
                parsed_data = extract_eml_contents(file_path)
            elif ext in ['.pdf', '.xlsx', '.xls', '.xlsm', '.csv', '.docx', '.png', '.jpg', '.jpeg', '.bmp', '.webp', '.txt']:
                parsed_data = extract_standalone_file_contents(file_path)
            else:
                print(f"Skipping unsupported file type: {filename}")
                continue
                
            # Execute LLM Extraction
            print("Extracting via LLM...")
            extracted_soa, _usage = call_llm_for_extraction(parsed_data)
            
            # Map results to Standard Template and append to master list
            for item in extracted_soa.line_items:
                # Format: ACTUAL_EMAIL_NAME/source_document
                # Example: "FW SOA.msg/Baker Ing.pdf" or "FW SOA.msg/Email Body"
                
                doc_name = item.source_document if item.source_document else "Email Body"
                source_formatted = f"{filename}/{doc_name}"
                
                all_records.append({
                    "Date": item.invoice_date,
                    "Invoice Number": item.invoice_number,
                    "Amount( Original )": item.invoice_value,
                    "Open amount": item.open_amount,
                    "Invoice Status": "",
                    "Pending with": "",
                    "Source File": source_formatted 
                })
                
            print(f"-> Extracted {len(extracted_soa.line_items)} line items from {filename}")
            
        except Exception as e:
            print(f"-> Error processing {filename}: {str(e)}")
            
    # Compile Master Excel Output
    if all_records:
        columns = ["Date", "Invoice Number", "Amount( Original )", "Open amount", "Invoice Status", "Pending with", "Source File"]
        df = pd.DataFrame(all_records, columns=columns)
        
        with pd.ExcelWriter(output_excel_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='SOA_Status', index=False)
            
        print(f"\nSUCCESS! Exported a total of {len(all_records)} line items to {output_excel_path} in sheet 'SOA_Status'.")
    else:
        print("\nNo records were extracted from any files.")

if __name__ == "__main__":
    folder_to_process = r"C:\Users\SohaibNM\OneDrive - Otis Global Tenant\Documents\Projects\VSR_UK\Data"
    output_path = r"C:\Users\SohaibNM\OneDrive - Otis Global Tenant\Documents\Projects\VSR_UK\Master_Output_SOA_Status.xlsx"
    
    batch_process_folder(folder_to_process, output_path)
